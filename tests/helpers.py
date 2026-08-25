"""In-memory docx builders for synthetic tests."""

from io import BytesIO

from docx import Document
from lxml import etree

from docxcast._ooxml import qn


def append_to_body(doc: Document, el) -> None:
    body = doc.element.body
    sect = body.find(qn('w:sectPr'))
    if sect is None:
        body.append(el)
    else:
        sect.addprevious(el)


def _as_list(el) -> list:
    return el if isinstance(el, list) else [el]


def make_sdt(alias=None, tag=None, *, pr_extra=None, content=None):
    sdt = etree.Element(qn('w:sdt'))
    pr = etree.SubElement(sdt, qn('w:sdtPr'))
    if alias is not None:
        el = etree.SubElement(pr, qn('w:alias'))
        el.set(qn('w:val'), alias)
    if tag is not None:
        el = etree.SubElement(pr, qn('w:tag'))
        el.set(qn('w:val'), tag)
    if pr_extra is not None:
        pr.extend(_as_list(pr_extra))
    c = etree.SubElement(sdt, qn('w:sdtContent'))
    if content is not None:
        c.extend(_as_list(content))
    return sdt


def run(text: str):
    r = etree.Element(qn('w:r'))
    t = etree.SubElement(r, qn('w:t'))
    t.text = text
    return r


def para(*children):
    p = etree.Element(qn('w:p'))
    for c in children:
        p.append(c)
    return p


def lock_pr(val: str = 'sdtLocked', lock_type=None):
    el = etree.Element(qn('w:lock'))
    el.set(qn('w:val'), val)
    if lock_type is not None:
        el.set(qn('w:lockType'), lock_type)
    return el


def dropdown_pr(items):
    dd = etree.Element(qn('w:dropDownList'))
    for display, value in items:
        li = etree.SubElement(dd, qn('w:listItem'))
        li.set(qn('w:displayText'), display)
        li.set(qn('w:value'), value)
    return dd


def date_pr(store: str = 'date', date_format: str | None = None):
    d = etree.Element(qn('w:date'))
    if date_format:
        el = etree.SubElement(d, qn('w:dateFormat'))
        el.set(qn('w:val'), date_format)
    el = etree.SubElement(d, qn('w:storeMappedDataAs'))
    el.set(qn('w:val'), store)
    return d


def checkbox_pr():
    cb = etree.Element(qn('w14:checkbox'))
    el = etree.SubElement(cb, qn('w14:checked'))
    el.set(qn('w14:val'), '0')
    for tag, val in ((qn('w14:checkedState'), '2612'), (qn('w14:uncheckedState'), '2610')):
        el = etree.SubElement(cb, tag)
        el.set(qn('w14:val'), val)
    return cb


def repeating_pr():
    return etree.Element(qn('w15:repeatingSection'))


def group_pr():
    return etree.Element(qn('w:group'))


def repeating_item_pr():
    return etree.Element(qn('w15:repeatingSectionItem'))


def docpartlist_pr():
    return etree.Element(qn('w:docPartList'))


def picture_pr():
    return etree.Element(qn('w:picture'))


def build_doc(*children) -> bytes:
    doc = Document()
    for c in children:
        append_to_body(doc, c)
    return save_bytes(doc)


def build_cell_doc(el, *, replace_paragraph: bool = False) -> bytes:
    """A one-cell table holding ``el``: appended to the cell's paragraph,
    or replacing it when the element is block-level."""
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    if replace_paragraph:
        cell.paragraphs[0]._p.getparent().replace(cell.paragraphs[0]._p, el)
    else:
        cell.paragraphs[0]._p.append(el)
    return save_bytes(doc)


def empty_template() -> bytes:
    return save_bytes(Document())


def full_text(doc: Document) -> str:
    return ''.join(t.text or '' for t in doc.element.body.iter(qn('w:t')))


def count_sdts(doc: Document) -> int:
    return len(list(doc.element.body.iter(qn('w:sdt'))))


def reopen(result) -> Document:
    return Document(BytesIO(result.document))


def save_bytes(doc: Document) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
